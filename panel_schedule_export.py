"""Export an ACIES panel-schedule workbook (many sheets) to organized text data.

Every worksheet that holds a panel schedule becomes one entry with its header
block, circuit list, connected-load totals, demand-load calculation and notes.
The result is written as Markdown (easy to read or paste into an AI chat) and
JSON (easy to process).

The parser is driven by the labels printed on the sheet (``VOLTAGE:``,
``CKT#``, ``DEMAND LOAD`` ...) instead of fixed cell addresses, so it tolerates
the small row/column differences between sheets. Workbooks are opened
read-only; ``.xls`` needs Microsoft Excel (same as the panel-schedule sync),
``.xlsx`` falls back to openpyxl.

Usage::

    python panel_schedule_export.py "260487 Panels.xls"
    python panel_schedule_export.py "260487 Panels.xls" -o C:/temp/out --include-hidden
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

try:
    from apps.ProjectManagement.panel_schedule_sync import (
        SUPPORTED_EXTENSIONS,
        PanelScheduleSyncError,
        _com_initialized,
        _normalize_matrix,
        _release_com_object,
    )
except ImportError:  # pragma: no cover - direct execution fallback
    from panel_schedule_sync import (  # type: ignore
        SUPPORTED_EXTENSIONS,
        PanelScheduleSyncError,
        _com_initialized,
        _normalize_matrix,
        _release_com_object,
    )

# Sheets that are never panels, matched case-insensitively by name.
SKIP_SHEET_NAMES = {"TEMPLATE"}
# Largest region read from a sheet; real schedules are ~70 rows x 26 columns.
MAX_READ_ROWS = 200
MAX_READ_COLUMNS = 40

LOAD_TYPE_NAMES = {
    "C": "Continuous",
    "D": "Dedicated",
    "G": "General receptacle",
    "K": "Kitchen equipment",
    "M": "Motor",
    "ML": "Largest motor",
}
STATUS_NAMES = {"N": "New", "E": "Existing"}

# Excel stores formula errors in Value2 as these negative integers.
_EXCEL_ERRORS = {
    -2146826281: "#DIV/0!",
    -2146826246: "#N/A",
    -2146826259: "#NAME?",
    -2146826288: "#NULL!",
    -2146826252: "#NUM!",
    -2146826265: "#REF!",
    -2146826273: "#VALUE!",
}

Matrix = List[List[Any]]


# --------------------------------------------------------------------------
# Cell helpers (all coordinates are 1-based like Excel)
# --------------------------------------------------------------------------


def _col_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _cell(matrix: Matrix, row: int, column: int) -> Any:
    if row < 1 or column < 1 or row > len(matrix):
        return None
    cells = matrix[row - 1]
    return cells[column - 1] if column <= len(cells) else None


def _number(value: Any) -> Optional[float]:
    """A numeric cell (or numeric text such as '20') as float, else None."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        if int(value) == value and int(value) in _EXCEL_ERRORS:
            return None
        return float(value)
    text = str(value).strip()
    if re.fullmatch(r"-?\d+(?:\.\d+)?", text):
        return float(text)
    return None


def _tidy_number(value: float) -> Any:
    rounded = round(value, 4)
    return int(rounded) if rounded == int(rounded) else rounded


def _text(value: Any) -> str:
    """Display text for a cell: whole numbers lose '.0', errors become '#REF!'."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        if int(value) == value and int(value) in _EXCEL_ERRORS:
            return _EXCEL_ERRORS[int(value)]
        return str(_tidy_number(float(value)))
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in str(value).splitlines()]
    return "\n".join(line for line in lines if line)


def _is_error(value: Any) -> bool:
    return _text(value) in _EXCEL_ERRORS.values()


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", _text(value)).strip().upper()


def _find_cells(matrix: Matrix, predicate, rows: Optional[range] = None) -> List[Tuple[int, int]]:
    found: List[Tuple[int, int]] = []
    for row in rows or range(1, len(matrix) + 1):
        for column in range(1, len(matrix[row - 1]) + 1):
            if predicate(_norm(_cell(matrix, row, column))):
                found.append((row, column))
    return found


def _first_number_right(matrix: Matrix, row: int, start_column: int) -> Optional[float]:
    for column in range(start_column, len(matrix[row - 1]) + 1):
        number = _number(_cell(matrix, row, column))
        if number is not None:
            return number
    return None


# --------------------------------------------------------------------------
# Single-sheet parsing
# --------------------------------------------------------------------------


def _split_title(title: str, sheet_name: str) -> Dict[str, str]:
    """'(N) PANEL 'F'\\nDETACHED GARAGE' -> status N, name F, description."""
    lines = [line.strip() for line in title.splitlines() if line.strip()]
    first = lines[0] if lines else ""
    description = " - ".join(lines[1:])
    status = ""
    match = re.match(r"\(\s*([A-Za-z])\s*\)\s*", first)
    if match:
        status = match.group(1).upper()
        first = first[match.end():]
    name_match = re.search(r"\b(?:PANEL|PNL)\s+['\"]?([^'\"]+?)['\"]?\s*$", first, re.IGNORECASE)
    name = name_match.group(1).strip() if name_match else (first or sheet_name)
    return {
        "name": name,
        "statusCode": status,
        "status": STATUS_NAMES.get(status, status),
        "description": description,
    }


def _header_fields(matrix: Matrix, header_row: int) -> Tuple[Dict[str, str], List[str]]:
    """Read 'LABEL: value' pairs above the circuit table, plus the NOTES: stack."""
    fields: Dict[str, str] = {}
    side_notes: List[str] = []
    labels = _find_cells(
        matrix, lambda text: text.endswith(":") and len(text) > 1, range(1, header_row)
    )
    for row, column in labels:
        label = _norm(_cell(matrix, row, column)).rstrip(":").strip()
        if label == "NOTES":
            for below in range(row + 1, header_row):
                value = _text(_cell(matrix, below, column))
                if value:
                    side_notes.append(value)
            continue
        for right in range(column + 1, len(matrix[row - 1]) + 1):
            value = _text(_cell(matrix, row, right))
            if _norm(value).endswith(":"):
                break
            if value:
                fields[label] = value
                break
    return fields, side_notes


def _locate_circuit_table(matrix: Matrix) -> Optional[Dict[str, Any]]:
    """Find the header row and map the left/right column groups by their labels."""
    for row in range(1, min(len(matrix), 15) + 1):
        ckt_columns = [c for c in range(1, len(matrix[row - 1]) + 1) if _norm(_cell(matrix, row, c)) == "CKT#"]
        if len(ckt_columns) < 2:
            continue
        phase_column = next(
            (c for c in range(1, len(matrix[row - 1]) + 1) if _norm(_cell(matrix, row, c)) == "PHASE"),
            None,
        )
        if phase_column is None:
            continue
        sub_row = row + 1
        sub_labels = {
            c: _norm(_cell(matrix, sub_row, c)) for c in range(1, len(matrix[sub_row - 1]) + 1)
        } if sub_row <= len(matrix) else {}
        has_sub_row = any(label in ("P", "T") for label in sub_labels.values())

        def group(low: int, high: int) -> Dict[str, int]:
            mapping: Dict[str, int] = {}
            for column in range(low, high):
                label = _norm(_cell(matrix, row, column))
                key = {
                    "CKT#": "ckt",
                    "NOTES": "notes",
                    "LOAD": "type",
                    "CB": "cb",
                    "LOAD DESCRIPTION": "desc",
                    "LOAD KVA": "kva",
                }.get(label)
                if key and key not in mapping:
                    mapping[key] = column
            cb = mapping.get("cb")
            if cb is not None:
                poles = [c for c in range(low, high) if sub_labels.get(c) == "P"]
                trips = [c for c in range(low, high) if sub_labels.get(c) == "T"]
                mapping["poles"] = poles[0] if poles else cb
                mapping["trip"] = trips[0] if trips else cb + 1
            return mapping

        left = group(1, phase_column)
        right = group(phase_column + 1, len(matrix[row - 1]) + 1)
        if "ckt" not in left or "ckt" not in right or "desc" not in left or "desc" not in right:
            continue
        return {
            "headerRow": row,
            "firstDataRow": row + (2 if has_sub_row else 1),
            "phaseColumn": phase_column,
            "left": left,
            "right": right,
        }
    return None


def _note_codes(raw: str) -> List[str]:
    return [token for token in re.split(r"[,/;\s]+", raw) if token]


def _parse_circuits(
    values: Matrix, formulas: Matrix, table: Dict[str, Any]
) -> Tuple[List[Dict[str, Any]], int, List[str]]:
    """Return (circuits, last circuit row, warnings). Multi-pole breakers span rows."""
    first_row = table["firstDataRow"]
    last_row = first_row - 1
    for row in range(first_row, len(values) + 1):
        if _norm(_cell(values, row, 1)).startswith("CONNECTED LOAD"):
            break
        last_row = row
    warnings: List[str] = []
    circuits: List[Dict[str, Any]] = []

    for side in ("left", "right"):
        cols = table[side]
        current: Optional[Dict[str, Any]] = None
        for row in range(first_row, last_row + 1):
            def read(key: str, source: Matrix = values) -> Any:
                return _cell(source, row, cols[key]) if key in cols else None

            label_value = read("ckt")
            label = _text(label_value)
            if not label:
                continue
            phase = _norm(_cell(values, row, table["phaseColumn"]))
            kva_value = read("kva")
            kva = _number(kva_value)
            if _is_error(kva_value):
                warnings.append(f"Row {row}: load kVA cell is {_text(kva_value)}")
            formula = _cell(formulas, row, cols["kva"]) if "kva" in cols else None
            basis = ""
            if isinstance(formula, str) and re.fullmatch(r"=[\d.\s*/+()-]+", formula.strip()):
                basis = formula.strip()[1:].replace(" ", "")
            poles = _number(read("poles"))
            trip = _number(read("trip"))

            continues = (
                poles is None
                and trip is None
                and current is not None
                and (current["poles"] or 0) > len(current["circuits"])
            )
            if continues:
                current["circuits"].append(_tidy_label(label_value))
                current["sheetRows"].append(row)
                current["phases"].append(phase)
                target = current
            else:
                description = _text(read("desc"))
                if description in ("---", "--", "-"):
                    description = ""
                code = _norm(read("type"))
                kind = (
                    "spare" if description.upper() == "SPARE"
                    else "space" if description.upper() == "SPACE"
                    else "load" if description else "empty"
                )
                notes = _text(read("notes"))
                target = {
                    "circuit": _tidy_label(label_value),
                    "circuits": [_tidy_label(label_value)],
                    "side": side,
                    "sheetRows": [row],
                    "phases": [phase],
                    "poles": int(poles) if poles is not None else None,
                    "breakerAmps": _tidy_number(trip) if trip is not None else None,
                    "kind": kind,
                    "description": description.replace("\n", " "),
                    "loadTypeCode": code,
                    "loadType": LOAD_TYPE_NAMES.get(code, ""),
                    "kva": None,
                    "kvaByPhase": {},
                    "kvaBasis": "",
                    "notes": notes,
                    "noteCodes": _note_codes(notes),
                }
                circuits.append(target)
                current = target
            if not target["kvaBasis"] and basis:
                target["kvaBasis"] = basis
            if kva is not None:
                target["kva"] = (target["kva"] or 0.0) + kva
                if phase:
                    by_phase = target["kvaByPhase"]
                    by_phase[phase] = by_phase.get(phase, 0.0) + kva

    for circuit in circuits:
        if circuit["kva"] is not None:
            circuit["kva"] = _tidy_number(circuit["kva"])
        circuit["kvaByPhase"] = {k: _tidy_number(v) for k, v in circuit["kvaByPhase"].items()}
        circuit["phases"] = [p for p in circuit["phases"] if p]

    def order(circuit: Dict[str, Any]) -> Tuple[int, float, int]:
        number = _number(circuit["circuit"])
        return (0, number, 0) if number is not None else (1, 0.0, circuit["sheetRows"][0])

    circuits.sort(key=order)
    return circuits, last_row, warnings


def _tidy_label(value: Any) -> Any:
    number = _number(value)
    return int(number) if number is not None and number == int(number) else _text(value)


def _parse_connected_load(values: Matrix, after_row: int) -> Dict[str, Any]:
    result: Dict[str, Any] = {}
    start = next(
        (r for r in range(after_row, len(values) + 1)
         if _norm(_cell(values, r, 1)).startswith("CONNECTED LOAD")),
        None,
    )
    if start is None:
        return result
    for row in range(start + 1, min(start + 8, len(values)) + 1):
        label = _norm(_cell(values, row, 1))
        key = {"PHASE A": "A", "PHASE B": "B", "PHASE C": "C", "TOTAL": "total"}.get(label)
        if not key:
            if label == "NOTES":
                break
            continue
        number = _first_number_right(values, row, 2)
        if number is not None:
            result[key] = _tidy_number(number)
    return result


def _parse_demand(values: Matrix, after_row: int) -> List[Dict[str, Any]]:
    """Each 'DEMAND LOAD' header row starts a block (main panel or SUB PANEL n)."""
    blocks: List[Dict[str, Any]] = []
    headers = _find_cells(values, lambda text: text == "DEMAND LOAD", range(after_row, len(values) + 1))
    for index, (header_row, label_column) in enumerate(headers):
        row_text = {c: _norm(_cell(values, header_row, c)) for c in range(1, len(values[header_row - 1]) + 1)}
        subtotal_col = next((c for c, t in row_text.items() if t == "SUBTOTAL"), None)
        factor_col = next((c for c, t in row_text.items() if t.startswith("NEC DEMAND")), None)
        if subtotal_col is None:
            continue
        above = _text(_cell(values, header_row - 1, label_column))
        name = above if above and not _norm(above).startswith("DEMAND LOAD") else "MAIN"
        limit = headers[index + 1][0] - 1 if index + 1 < len(headers) else len(values)
        lines: List[Dict[str, Any]] = []
        total_kva = amps = None
        amps_basis = ""
        for row in range(header_row + 1, limit + 1):
            label = _text(_cell(values, row, label_column))
            if factor_col is not None:
                factor_text = _text(_cell(values, row, factor_col))
                result = _first_number_right(values, row, factor_col + 1)
                if _norm(factor_text).startswith("TOTAL DEMAND") and result is not None:
                    total_kva = _tidy_number(result)
                elif _norm(factor_text).startswith("AMPS @") and result is not None:
                    amps = _tidy_number(result)
                    amps_basis = factor_text
            if not label or _norm(label).startswith("SUB PANEL"):
                continue
            subtotal = _number(_cell(values, row, subtotal_col))
            line: Dict[str, Any] = {"label": label, "subtotal": _tidy_number(subtotal) if subtotal is not None else None}
            if factor_col is not None:
                factor_cell = _cell(values, row, factor_col)
                factor_number = _number(factor_cell)
                factor = _tidy_number(factor_number) if factor_number is not None else _text(factor_cell)
                demand = _first_number_right(values, row, factor_col + 1)
                if _norm(factor).startswith(("TOTAL DEMAND", "AMPS @")):
                    factor, demand = "", None
                line["factor"] = factor
                line["demand"] = _tidy_number(demand) if demand is not None else None
            lines.append(line)
        if not any(line["subtotal"] is not None for line in lines):
            continue  # an unused sub-panel placeholder
        block: Dict[str, Any] = {"name": name, "lines": lines}
        if total_kva is not None:
            block["totalDemandKva"] = total_kva
        if amps is not None:
            block["demandAmps"] = amps
            block["demandAmpsBasis"] = amps_basis
        blocks.append(block)
    return blocks


def _parse_footer_notes(values: Matrix, after_row: int) -> List[str]:
    start = next(
        (r for r in range(after_row, len(values) + 1) if _norm(_cell(values, r, 1)) == "NOTES"),
        None,
    )
    if start is None:
        return []
    notes: List[str] = []
    blank_run = 0
    for row in range(start + 1, len(values) + 1):
        text = _text(_cell(values, row, 1)).replace("\n", " ")
        if not text:
            blank_run += 1
            if blank_run >= 2:
                break
            continue
        blank_run = 0
        notes.append(text)
    return notes


def _amps(text: str) -> Optional[int]:
    match = re.search(r"(\d+)\s*A\b", text, re.IGNORECASE)
    return int(match.group(1)) if match else None


def parse_panel_matrix(sheet_name: str, values: Matrix, formulas: Optional[Matrix] = None) -> Optional[Dict[str, Any]]:
    """Parse one sheet; None when it has no panel circuit table."""
    formulas = formulas if formulas is not None else values
    table = _locate_circuit_table(values)
    if table is None:
        return None
    title = _text(_cell(values, 3, 1)) or next(
        (_text(_cell(values, r, 1)) for r in range(1, table["headerRow"]) if _text(_cell(values, r, 1))), ""
    )
    header = _split_title(title, sheet_name)
    fields, side_notes = _header_fields(values, table["headerRow"])
    circuits, last_row, warnings = _parse_circuits(values, formulas, table)
    connected = _parse_connected_load(values, last_row)
    demand = _parse_demand(values, last_row)

    main_text = fields.get("MAIN REQUIREMENT", "")
    main_type = "MLO" if "MLO" in main_text.upper() else ("MCB" if "MCB" in main_text.upper() else "")
    phases_text = fields.get("PHASE", "")
    wire_text = fields.get("WIRE", "")
    panel: Dict[str, Any] = {
        "sheet": sheet_name,
        "name": header["name"],
        "status": header["status"],
        "description": header["description"],
        "voltage": fields.get("VOLTAGE", ""),
        "phase": _tidy_number(_number(phases_text)) if _number(phases_text) is not None else phases_text,
        "wire": _tidy_number(_number(wire_text)) if _number(wire_text) is not None else wire_text,
        "busRating": fields.get("BUS RATING", ""),
        "busAmps": _amps(fields.get("BUS RATING", "")),
        "mainRequirement": main_text,
        "mainType": main_type,
        "mainAmps": _amps(main_text),
        "enclosure": fields.get("TYPE", ""),
        "mounting": fields.get("MOUNTING", ""),
        "aicRating": fields.get("AIC RATING", ""),
        "condition": fields.get("CONDITION", ""),
        "fedFrom": side_notes[0] if side_notes else "",
        "headerNotes": side_notes,
        "circuits": circuits,
        "connectedLoadKva": connected,
        "demand": demand,
        "notes": _parse_footer_notes(values, last_row),
        "warnings": warnings,
    }
    panel["warnings"].extend(_check_connected_load(circuits, connected))
    return panel


def _check_connected_load(circuits: List[Dict[str, Any]], connected: Dict[str, Any]) -> List[str]:
    """Compare the sheet's printed per-phase totals with the circuit rows."""
    sums = {"A": 0.0, "B": 0.0, "C": 0.0}
    for circuit in circuits:
        for phase, kva in circuit["kvaByPhase"].items():
            if phase in sums:
                sums[phase] += kva
    warnings = []
    for phase, total in sums.items():
        printed = connected.get(phase)
        if printed is not None and abs(printed - total) > 0.05:
            warnings.append(
                f"Printed connected load for phase {phase} is {printed} kVA but the circuit rows add up to {round(total, 3)} kVA"
            )
    return warnings


def parse_notes_matrix(values: Matrix) -> Optional[Dict[str, Any]]:
    """The 'PANEL SCHEDULE NOTES' legend sheet: general notes + keyed notes."""
    if not _find_cells(values, lambda text: text == "PANEL SCHEDULE NOTES", range(1, min(len(values), 3) + 1)):
        return None
    general: List[str] = []
    keyed: Dict[str, str] = {}
    mode = ""
    for row in range(1, len(values) + 1):
        first, second = _text(_cell(values, row, 1)), _text(_cell(values, row, 2))
        if _norm(first) == "GENERAL NOTES:":
            mode = "general"
        elif _norm(first) == "CIRCUIT KEYED NOTES:":
            mode = "keyed"
        elif mode == "general" and second:
            general.append(second)
        elif mode == "keyed" and first and second:
            keyed[first] = second
    return {"generalNotes": general, "keyedNotes": keyed}


# --------------------------------------------------------------------------
# Workbook reading
# --------------------------------------------------------------------------


def _read_sheets_with_excel(path: str) -> List[Dict[str, Any]]:
    try:
        import win32com.client
    except Exception as exc:  # pragma: no cover - depends on Windows runtime
        raise PanelScheduleSyncError(
            "Reading .xls panel schedules requires Microsoft Excel and pywin32."
        ) from exc
    sheets: List[Dict[str, Any]] = []
    excel = workbook = None
    with _com_initialized():
        try:
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            excel.AskToUpdateLinks = False
            try:
                excel.AutomationSecurity = 3
            except Exception:
                pass
            workbook = excel.Workbooks.Open(
                os.path.abspath(path), UpdateLinks=0, ReadOnly=True,
                IgnoreReadOnlyRecommended=True, AddToMru=False,
            )
            for index in range(1, int(workbook.Worksheets.Count) + 1):
                worksheet = used = region = None
                try:
                    worksheet = workbook.Worksheets.Item(index)
                    used = worksheet.UsedRange
                    rows = min(MAX_READ_ROWS, int(used.Row) + int(used.Rows.Count) - 1)
                    columns = min(MAX_READ_COLUMNS, int(used.Column) + int(used.Columns.Count) - 1)
                    region = worksheet.Range(worksheet.Cells(1, 1), worksheet.Cells(rows, columns))
                    sheets.append({
                        "name": str(worksheet.Name),
                        "visible": int(worksheet.Visible) == -1,
                        "values": _normalize_matrix(region.Value2, rows, columns),
                        "formulas": _normalize_matrix(region.Formula, rows, columns),
                    })
                finally:
                    _release_com_object(region)
                    _release_com_object(used)
                    _release_com_object(worksheet)
        except PanelScheduleSyncError:
            raise
        except Exception as exc:
            raise PanelScheduleSyncError(f"Could not read panel schedule workbook: {exc}") from exc
        finally:
            if workbook is not None:
                try:
                    workbook.Close(False)
                except Exception:
                    pass
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass
            _release_com_object(workbook)
            _release_com_object(excel)
    return sheets


def _read_sheets_with_openpyxl(path: str) -> List[Dict[str, Any]]:
    import openpyxl

    formula_book = openpyxl.load_workbook(path, data_only=False)
    data_book = openpyxl.load_workbook(path, data_only=True)
    sheets: List[Dict[str, Any]] = []
    try:
        for name in formula_book.sheetnames:
            formula_ws, data_ws = formula_book[name], data_book[name]
            rows = min(MAX_READ_ROWS, int(formula_ws.max_row or 1))
            columns = min(MAX_READ_COLUMNS, int(formula_ws.max_column or 1))
            sheets.append({
                "name": name,
                "visible": formula_ws.sheet_state == "visible",
                "values": [[data_ws.cell(r, c).value for c in range(1, columns + 1)] for r in range(1, rows + 1)],
                "formulas": [[formula_ws.cell(r, c).value for c in range(1, columns + 1)] for r in range(1, rows + 1)],
            })
    finally:
        formula_book.close()
        data_book.close()
    return sheets


def read_workbook_sheets(path: str) -> List[Dict[str, Any]]:
    """Every sheet as {name, visible, values, formulas} matrices (read-only)."""
    workbook_path = os.path.abspath(str(path or "").strip())
    if not os.path.isfile(workbook_path):
        raise FileNotFoundError(f"Panel workbook not found at: {workbook_path}")
    extension = Path(workbook_path).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise PanelScheduleSyncError("Panel schedule must be an .xls or .xlsx workbook.")
    if extension == ".xls":
        return _read_sheets_with_excel(workbook_path)
    return _read_sheets_with_openpyxl(workbook_path)


# --------------------------------------------------------------------------
# Whole-workbook export
# --------------------------------------------------------------------------


def export_sheets(
    sheets: Iterable[Dict[str, Any]], source: str = "", include_hidden: bool = False
) -> Dict[str, Any]:
    """Parse already-read sheets into the export structure."""
    panels: List[Dict[str, Any]] = []
    skipped: List[Dict[str, str]] = []
    legend: Optional[Dict[str, Any]] = None
    for sheet in sheets:
        name, values = sheet["name"], sheet["values"]
        notes = parse_notes_matrix(values)
        if notes is not None:
            if legend is None or sheet.get("visible", True):
                legend = notes
            skipped.append({"sheet": name, "reason": "notes legend (read separately)"})
            continue
        if name.strip().upper() in SKIP_SHEET_NAMES:
            skipped.append({"sheet": name, "reason": "blank template"})
            continue
        if not sheet.get("visible", True) and not include_hidden:
            skipped.append({"sheet": name, "reason": "hidden (use --include-hidden)"})
            continue
        panel = parse_panel_matrix(name, values, sheet.get("formulas"))
        if panel is None:
            skipped.append({"sheet": name, "reason": "no panel circuit table found"})
            continue
        panel["hidden"] = not sheet.get("visible", True)
        panels.append(panel)
    return {
        "source": source,
        "exportedAt": datetime.now().isoformat(timespec="seconds"),
        "panelCount": len(panels),
        "panels": panels,
        "notesLegend": legend or {"generalNotes": [], "keyedNotes": {}},
        "skippedSheets": skipped,
    }


def export_workbook(path: str, include_hidden: bool = False) -> Dict[str, Any]:
    return export_sheets(read_workbook_sheets(path), source=os.path.abspath(path), include_hidden=include_hidden)


# --------------------------------------------------------------------------
# Markdown rendering
# --------------------------------------------------------------------------


def _md(text: Any) -> str:
    return str(text).replace("|", "\\|").replace("\n", " ")


def _kva(value: Any) -> str:
    return "" if value is None else f"{float(value):.2f}"


def _factor_text(factor: Any) -> str:
    """1.25 -> '125%'; descriptive text is shown as printed."""
    return f"{factor * 100:g}%" if isinstance(factor, (int, float)) else str(factor)


def _join(parts: Iterable[str], sep: str = " · ") -> str:
    return sep.join(part for part in parts if part)


def render_panel_markdown(panel: Dict[str, Any]) -> str:
    """One panel as Markdown: header, circuits, demand calculation, notes, checks."""
    out: List[str] = []
    out.append(f"## Panel {panel['name']}" + (f" - {panel['status']}" if panel["status"] else ""))
    out.append("")
    if panel["description"]:
        out.append(f"**{panel['description']}**  ")
    out.append(f"Sheet: `{panel['sheet']}`")
    out.append("")
    phase = f"{panel['phase']}-phase" if panel["phase"] != "" else ""
    wire = f"{panel['wire']}-wire" if panel["wire"] != "" else ""
    out.append(f"- Voltage: {_join([panel['voltage'], phase, wire])}")
    out.append(f"- Bus: {panel['busRating']} · Main: {panel['mainRequirement'].strip() or 'n/a'}")
    out.append(
        "- " + _join([
            f"Mounting: {panel['mounting']}", f"Enclosure: {panel['enclosure']}",
            f"AIC: {panel['aicRating']}", f"Condition: {panel['condition']}",
        ])
    )
    if panel["headerNotes"]:
        out.append(f"- Feed: {_join(panel['headerNotes'])}")
    out.append("")

    out.append("### Circuits")
    out.append("")
    out.append("| Ckt | Phase | Breaker | Load | Description | kVA | Notes |")
    out.append("|---|---|---|---|---|---|---|")
    for circuit in panel["circuits"]:
        ckt = "/".join(dict.fromkeys(str(c) for c in circuit["circuits"]))
        breaker = (
            f"{circuit['poles'] or '?'}P / {circuit['breakerAmps']}A"
            if circuit["breakerAmps"] is not None else ""
        )
        description = circuit["description"] or ("" if circuit["kind"] == "empty" else circuit["kind"].upper())
        out.append(
            f"| {_md(ckt)} | {'/'.join(circuit['phases'])} | {breaker} | {circuit['loadTypeCode']} "
            f"| {_md(description)} | {_kva(circuit['kva'])} | {_md(circuit['notes'])} |"
        )
    out.append("")

    connected = panel["connectedLoadKva"]
    if connected:
        out.append(
            "Connected load: " + _join(
                [f"{k} {_kva(connected[k])}" for k in ("A", "B", "C") if k in connected]
                + ([f"total {_kva(connected['total'])} kVA"] if "total" in connected else [])
            )
        )
        out.append("")

    for block in panel["demand"]:
        out.append(f"### Demand load - {block['name']}")
        out.append("")
        has_factor = any("factor" in line for line in block["lines"])
        if has_factor:
            out.append("| Load | Subtotal | Demand factor | Demand |")
            out.append("|---|---|---|---|")
        else:
            out.append("| Load | Subtotal |")
            out.append("|---|---|")
        for line in block["lines"]:
            subtotal = _kva(line["subtotal"])
            if has_factor:
                out.append(f"| {_md(line['label'])} | {subtotal} | {_md(_factor_text(line.get('factor', '')))} | {_kva(line.get('demand'))} |")
            else:
                out.append(f"| {_md(line['label'])} | {subtotal} |")
        out.append("")
        if "totalDemandKva" in block:
            out.append(f"Total demand: {block['totalDemandKva']:.2f} kVA")
        if "demandAmps" in block:
            out.append(f"{block['demandAmpsBasis']}: {block['demandAmps']:.1f} A")
        if "totalDemandKva" in block or "demandAmps" in block:
            out.append("")

    if panel["notes"]:
        out.append("### Notes")
        out.append("")
        out.extend(f"{note}  " for note in panel["notes"])
        out.append("")
    if panel["warnings"]:
        out.append("### Check")
        out.append("")
        out.extend(f"- {warning}" for warning in panel["warnings"])
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def render_legend_markdown(legend: Dict[str, Any]) -> str:
    """The workbook notes legend as Markdown; empty when there is none."""
    if not (legend["generalNotes"] or legend["keyedNotes"]):
        return ""
    out: List[str] = []
    out.append("## Notes legend")
    out.append("")
    out.extend(f"- {note}" for note in legend["generalNotes"])
    if legend["keyedNotes"]:
        out.append("")
        out.append("Circuit keyed notes (codes in each circuit's Notes column; numbers refer to the panel's own notes):")
        out.append("")
        out.extend(f"- `{key}`: {text}" for key, text in legend["keyedNotes"].items())
    out.append("")
    return "\n".join(out).rstrip() + "\n"


def render_markdown(data: Dict[str, Any]) -> str:
    out: List[str] = []
    out.append(f"# Panel schedules - {Path(data['source']).name or 'workbook'}")
    out.append("")
    out.append(f"Source: `{data['source']}`  ")
    out.append(f"Exported: {data['exportedAt']} · {data['panelCount']} panel(s)")
    out.append("")
    out.append("## Contents")
    out.append("")
    for panel in data["panels"]:
        label = _join([f"Panel {panel['name']}", panel["status"], panel["description"]])
        out.append(f"- {label} (sheet `{panel['sheet']}`){' - hidden sheet' if panel.get('hidden') else ''}")
    out.append("")

    for panel in data["panels"]:
        out.append(render_panel_markdown(panel))
    legend = render_legend_markdown(data["notesLegend"])
    if legend:
        out.append(legend)

    if data["skippedSheets"]:
        out.append("## Skipped sheets")
        out.append("")
        out.extend(f"- `{item['sheet']}`: {item['reason']}" for item in data["skippedSheets"])
        out.append("")
    return "\n".join(out)


# --------------------------------------------------------------------------
# App entry point
# --------------------------------------------------------------------------


def _panel_summary(panel: Dict[str, Any]) -> str:
    loads = sum(1 for c in panel["circuits"] if c["kind"] == "load")
    spares = sum(1 for c in panel["circuits"] if c["kind"] in ("spare", "space"))
    total = panel["connectedLoadKva"].get("total")
    return _join([
        panel["voltage"],
        f"{panel['busAmps']}A bus" if panel["busAmps"] else "",
        f"{loads} load{'s' if loads != 1 else ''}",
        f"{spares} spare/space" if spares else "",
        f"{total:.1f} kVA" if total is not None else "",
    ])


def export_files(paths: Iterable[str], include_hidden: bool = False) -> List[Dict[str, Any]]:
    """Export several workbooks for the Panel Schedule Export dialog.

    One entry per workbook; a workbook that cannot be read carries an ``error``
    instead of failing the whole batch.
    """
    files: List[Dict[str, Any]] = []
    for path in paths:
        entry: Dict[str, Any] = {"path": str(path), "name": Path(str(path)).name}
        try:
            data = export_workbook(str(path), include_hidden=include_hidden)
        except Exception as exc:
            entry["error"] = str(exc) or exc.__class__.__name__
            files.append(entry)
            continue
        entry.update({
            "exportedAt": data["exportedAt"],
            "panels": [
                {
                    "key": f"{index}:{panel['sheet']}",
                    "name": panel["name"],
                    "sheet": panel["sheet"],
                    "status": panel["status"],
                    "description": panel["description"],
                    "hidden": bool(panel.get("hidden")),
                    "warningCount": len(panel["warnings"]),
                    "summary": _panel_summary(panel),
                    "markdown": render_panel_markdown(panel),
                    "data": panel,
                }
                for index, panel in enumerate(data["panels"])
            ],
            "legendMarkdown": render_legend_markdown(data["notesLegend"]),
            "legend": data["notesLegend"],
            "skippedSheets": data["skippedSheets"],
        })
        files.append(entry)
    return files


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def _safe_stem(path: str) -> str:
    return re.sub(r'[<>:"/\\|?*]+', "_", Path(path).stem).strip() or "panels"


def write_export(data: Dict[str, Any], out_dir: str, formats: Tuple[str, ...] = ("md", "json")) -> List[str]:
    os.makedirs(out_dir, exist_ok=True)
    stem = f"{_safe_stem(data['source'])} panel schedules"
    written: List[str] = []
    if "md" in formats:
        target = os.path.join(out_dir, stem + ".md")
        Path(target).write_text(render_markdown(data), encoding="utf-8", newline="\n")
        written.append(target)
    if "json" in formats:
        target = os.path.join(out_dir, stem + ".json")
        Path(target).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8", newline="\n")
        written.append(target)
    return written


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Export panel schedule sheets to Markdown/JSON text data.")
    parser.add_argument("workbook", help="Panel schedule workbook (.xls or .xlsx)")
    parser.add_argument("-o", "--out", help="Output folder (default: next to the workbook)")
    parser.add_argument("--include-hidden", action="store_true", help="Also export hidden sheets")
    parser.add_argument("--format", choices=("md", "json", "both"), default="both")
    args = parser.parse_args(argv)
    try:
        data = export_workbook(args.workbook, include_hidden=args.include_hidden)
    except (FileNotFoundError, PanelScheduleSyncError) as exc:
        print(f"Error: {exc}", file=sys.stderr)
        return 1
    formats = ("md", "json") if args.format == "both" else (args.format,)
    out_dir = args.out or os.path.dirname(os.path.abspath(args.workbook))
    for path in write_export(data, out_dir, formats):
        print(path)
    print(f"{data['panelCount']} panel(s) exported, {len(data['skippedSheets'])} sheet(s) skipped.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
