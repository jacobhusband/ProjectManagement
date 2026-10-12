"""
Unit tests for apps/ProjectManagement/panel_schedule_export.py
"""

import json
import os
import re
import tempfile
import unittest
from pathlib import Path

import openpyxl

import panel_schedule_export as export


def column_number(letters):
    value = 0
    for char in letters:
        value = value * 26 + ord(char) - 64
    return value


def grid(cells, rows=60, columns=26):
    matrix = [[None] * columns for _ in range(rows)]
    for address, value in cells.items():
        match = re.fullmatch(r"([A-Z]+)(\d+)", address)
        matrix[int(match.group(2)) - 1][column_number(match.group(1)) - 1] = value
    return matrix


HEADER = {
    "F2": "VOLTAGE:", "G2": "120/208V", "I2": "WIRE:", "K2": 4.0, "M2": "MOUNTING:", "N2": "SURFACE", "P2": "NOTES:",
    "A3": "(N) PANEL 'F'\nDETACHED GARAGE", "F3": "BUS RATING:", "G3": "200A - SEE NOTE #1",
    "I3": "PHASE:", "K3": 3.0, "M3": "AIC RATING:", "N3": "EXISTING", "P3": "FED FROM MAIN",
    "F4": "MAIN REQUIREMENT:", "G4": " MLO", "I4": "TYPE:", "K4": "NEMA 1R", "M4": "CONDITION:",
    "N4": "EXISTING", "P4": "DISTRIBUTION",
    "A6": "CKT#", "B6": "NOTES", "C6": "LOAD", "D6": "CB", "F6": "   LOAD DESCRIPTION", "I6": "LOAD KVA",
    "J6": "PHASE", "K6": "LOAD KVA", "L6": "LOAD DESCRIPTION", "O6": "CB", "Q6": "LOAD", "R6": "NOTES", "S6": "CKT#",
    "D7": "P", "E7": "T", "O7": "P", "P7": "T",
}


def panel_cells():
    cells = dict(HEADER)
    rows = {
        # row: (left ckt, notes, type, poles, trip, desc, kva, phase, right kva, desc, poles, trip, type, notes, ckt)
        8: (1, "2,AF", "G", 1, 20, "RECEPTACLES", 0.72, "A", 0.4, "INTERIOR LIGHTING", 1, "20", "C", "2,AF", 2),
        9: (3, "2,AF", "G", 1, 20, "RECEPTACLES", 0.54, "B", 0.1, "EXTERIOR LIGHTING", 1, "20", "C", "2,AF", 4),
        10: (5, "2", "G", 1, 20, "RECEPTACLES", 0.54, "C", 3.328, "EV CHARGER", 2, 40, "D", "3,GFI", 6),
        11: (7, "3", "D", 2, 30, "DRYER", 3.0, "A", 3.328, "---", "-", "-", "D", "3,GFI", 8),
        12: (9, "3", "D", "-", "-", "---", 3.0, "B", 0.5, "SPARE", 1, 20, None, None, 10),
        13: (11, None, None, 1, 20, "SPACE", None, "C", None, "SPARE", 1, 20, None, None, 12),
    }
    for row, (lc, ln, lt, lp, ltr, ld, lk, ph, rk, rd, rp, rtr, rt, rn, rc) in rows.items():
        for address, value in {
            f"A{row}": lc, f"B{row}": ln, f"C{row}": lt, f"D{row}": lp, f"E{row}": ltr, f"F{row}": ld,
            f"I{row}": lk, f"J{row}": ph, f"K{row}": rk, f"L{row}": rd, f"O{row}": rp, f"P{row}": rtr,
            f"Q{row}": rt, f"R{row}": rn, f"S{row}": rc,
        }.items():
            if value is not None:
                cells[address] = value
    cells.update({
        "A14": "CONNECTED LOAD:",
        "A15": "PHASE A", "F15": 7.448, "G15": "KVA",
        "A16": "PHASE B", "F16": 4.14, "G16": "KVA",
        "A17": "PHASE C", "F17": 3.868, "G17": "KVA",
        "A18": "TOTAL", "F18": 15.456,
        "I14": "DEMAND LOAD CALCULATION",
        "I15": "DEMAND LOAD", "M15": "SUBTOTAL", "N15": "NEC DEMAND FACTOR",
        "I16": "CONTINUOUS LOAD (C)", "M16": 0.5, "N16": 1.25, "Q16": 0.625,
        "I17": "DEDICATED LOAD (D)", "M17": 6.0, "N17": 1.0, "Q17": 6.0,
        "I18": "NUMBER OF KITCHEN EQUIP.", "M18": 0.0, "N18": "TOTAL DEMAND KVA", "Q18": 6.625,
        "N19": "AMPS @ 120/208V, 3 PHASE, 4 WIRE", "Q19": 18.4,
        "I21": "SUB PANEL 1", "I22": "DEMAND LOAD", "M22": "SUBTOTAL", "I23": "CONTINUOUS LOAD (C)",
        "I31": "SUB PANEL A1", "I32": "DEMAND LOAD", "M32": "SUBTOTAL", "I33": "CONTINUOUS LOAD (C)", "M33": 13.795,
        "A20": "NOTES", "A21": "1. FIRST NOTE.", "A22": "2. SECOND NOTE.",
    })
    return cells


class TestParsePanelMatrix(unittest.TestCase):
    def setUp(self):
        self.panel = export.parse_panel_matrix("F", grid(panel_cells()))
        self.by_ckt = {c["circuit"]: c for c in self.panel["circuits"]}

    def test_header_fields(self):
        panel = self.panel
        self.assertEqual(panel["name"], "F")
        self.assertEqual(panel["status"], "New")
        self.assertEqual(panel["description"], "DETACHED GARAGE")
        self.assertEqual((panel["voltage"], panel["phase"], panel["wire"]), ("120/208V", 3, 4))
        self.assertEqual((panel["busAmps"], panel["mainType"], panel["mainAmps"]), (200, "MLO", None))
        self.assertEqual((panel["mounting"], panel["enclosure"], panel["aicRating"]), ("SURFACE", "NEMA 1R", "EXISTING"))
        self.assertEqual(panel["headerNotes"], ["FED FROM MAIN", "DISTRIBUTION"])
        self.assertEqual(panel["fedFrom"], "FED FROM MAIN")

    def test_multi_pole_breaker_spans_rows_on_either_side(self):
        dryer = self.by_ckt[7]
        self.assertEqual(dryer["circuits"], [7, 9])
        self.assertEqual((dryer["poles"], dryer["breakerAmps"], dryer["kva"]), (2, 30, 6.0))
        self.assertEqual(dryer["phases"], ["A", "B"])
        ev = self.by_ckt[6]
        self.assertEqual((ev["circuits"], ev["poles"], ev["breakerAmps"]), ([6, 8], 2, 40))
        self.assertEqual(ev["kva"], 6.656)
        self.assertNotIn(9, self.by_ckt)
        self.assertNotIn(8, self.by_ckt)

    def test_string_trip_and_kinds(self):
        self.assertEqual(self.by_ckt[2]["breakerAmps"], 20)
        self.assertEqual(self.by_ckt[10]["kind"], "spare")
        self.assertEqual(self.by_ckt[11]["kind"], "space")
        self.assertEqual(self.by_ckt[1]["loadType"], "General receptacle")
        self.assertEqual(self.by_ckt[1]["noteCodes"], ["2", "AF"])

    def test_circuits_sorted_by_number(self):
        numbers = [c["circuit"] for c in self.panel["circuits"]]
        self.assertEqual(numbers, sorted(numbers))

    def test_connected_load_and_demand(self):
        self.assertEqual(self.panel["connectedLoadKva"], {"A": 7.448, "B": 4.14, "C": 3.868, "total": 15.456})
        self.assertEqual(self.panel["warnings"], [])
        main = self.panel["demand"][0]
        self.assertEqual(main["name"], "MAIN")
        self.assertEqual(main["totalDemandKva"], 6.625)
        self.assertEqual(main["demandAmps"], 18.4)
        self.assertEqual(main["lines"][0]["factor"], 1.25)
        self.assertEqual(main["lines"][0]["demand"], 0.625)

    def test_unused_sub_panel_blocks_are_dropped_and_used_ones_kept(self):
        self.assertEqual([block["name"] for block in self.panel["demand"]], ["MAIN", "SUB PANEL A1"])

    def test_footer_notes(self):
        self.assertEqual(self.panel["notes"], ["1. FIRST NOTE.", "2. SECOND NOTE."])

    def test_connected_load_mismatch_is_flagged(self):
        cells = panel_cells()
        cells["F15"] = 9.0
        panel = export.parse_panel_matrix("F", grid(cells))
        self.assertTrue(any("phase A" in warning for warning in panel["warnings"]))

    def test_excel_error_cells_are_flagged(self):
        cells = panel_cells()
        cells["K8"] = -2146826265
        panel = export.parse_panel_matrix("F", grid(cells))
        self.assertTrue(any("#REF!" in warning for warning in panel["warnings"]))

    def test_non_panel_sheet_returns_none(self):
        self.assertIsNone(export.parse_panel_matrix("Sheet1", grid({"A1": "hello"})))

    def test_subfeed_labels_and_three_pole_with_blank_continuations(self):
        cells = dict(HEADER)
        cells.update({
            "A8": 1, "D8": 3, "E8": 60, "F8": "FC-1", "I8": 15.789, "J8": "A", "S8": 2,
            "A9": 3, "I9": 15.789, "J9": "B", "S9": 4,
            "A10": 5, "I10": 15.789, "J10": "C", "S10": 6,
            "A11": "SF-L", "D11": 3, "E11": 200, "F11": "SUBPANEL", "I11": 1.0, "J11": "A", "S11": "SF-R",
            "A12": "SF-L", "D12": "-", "E12": "-", "F12": "---", "I12": 1.0, "J12": "B", "S12": "SF-R",
            "A13": "SF-L", "D13": "-", "E13": "-", "F13": "---", "I13": 1.0, "J13": "C", "S13": "SF-R",
            "A14": "CONNECTED LOAD:",
        })
        panel = export.parse_panel_matrix("X", grid(cells))
        numbers = [c["circuit"] for c in panel["circuits"]]
        self.assertEqual(numbers[0], 1)
        fc1 = panel["circuits"][0]
        self.assertEqual((fc1["circuits"], fc1["poles"], fc1["kva"]), ([1, 3, 5], 3, 47.367))
        feed = next(c for c in panel["circuits"] if c["circuit"] == "SF-L")
        self.assertEqual((feed["poles"], feed["breakerAmps"], feed["kva"]), (3, 200, 3.0))


class TestExportSheets(unittest.TestCase):
    def sheets(self):
        legend = grid({
            "A1": "PANEL SCHEDULE NOTES", "A2": "GENERAL NOTES:", "B3": "BALANCE PANELS",
            "A6": "CIRCUIT KEYED NOTES:", "A7": "AF", "B7": "ARC-FAULT BREAKER",
        })
        return [
            {"name": "F", "visible": True, "values": grid(panel_cells()), "formulas": grid(panel_cells())},
            {"name": "TEMPLATE", "visible": True, "values": grid(panel_cells()), "formulas": grid(panel_cells())},
            {"name": "Old", "visible": False, "values": grid(panel_cells()), "formulas": grid(panel_cells())},
            {"name": "Sheet1", "visible": True, "values": legend, "formulas": legend},
            {"name": "Junk", "visible": True, "values": grid({"A1": "x"}), "formulas": grid({"A1": "x"})},
        ]

    def test_hidden_template_notes_and_junk_sheets_are_skipped(self):
        data = export.export_sheets(self.sheets(), source="x.xls")
        self.assertEqual([p["sheet"] for p in data["panels"]], ["F"])
        reasons = {item["sheet"]: item["reason"] for item in data["skippedSheets"]}
        self.assertEqual(set(reasons), {"TEMPLATE", "Old", "Sheet1", "Junk"})
        self.assertIn("hidden", reasons["Old"])
        self.assertEqual(data["notesLegend"]["keyedNotes"], {"AF": "ARC-FAULT BREAKER"})
        self.assertEqual(data["notesLegend"]["generalNotes"], ["BALANCE PANELS"])

    def test_include_hidden(self):
        data = export.export_sheets(self.sheets(), source="x.xls", include_hidden=True)
        self.assertEqual([p["sheet"] for p in data["panels"]], ["F", "Old"])
        self.assertTrue(data["panels"][1]["hidden"])

    def test_markdown_contains_merged_circuits_and_legend(self):
        markdown = export.render_markdown(export.export_sheets(self.sheets(), source="x.xls"))
        self.assertIn("| 7/9 | A/B | 2P / 30A | D | DRYER | 6.00 |", markdown)
        self.assertIn("`AF`: ARC-FAULT BREAKER", markdown)
        self.assertIn("| CONTINUOUS LOAD (C) | 0.50 | 125% | 0.62 |", markdown)


class TestExportFiles(unittest.TestCase):
    def test_batch_payload_per_workbook_with_failures_isolated(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "Panels.xlsx")
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "F"
            for address, value in panel_cells().items():
                sheet[address] = value
            book.save(path)
            missing = os.path.join(folder, "missing.xlsx")

            files = export.export_files([path, missing])

        good, bad = files
        self.assertNotIn("error", good)
        panel = good["panels"][0]
        self.assertEqual((panel["name"], panel["sheet"], panel["warningCount"]), ("F", "F", 0))
        self.assertIn("## Panel F", panel["markdown"])
        self.assertIn("120/208V", panel["summary"])
        self.assertEqual(panel["data"]["circuits"][0]["description"], "RECEPTACLES")
        self.assertIn("not found", bad["error"])
        self.assertNotIn("panels", bad)


class TestAppWiring(unittest.TestCase):
    def test_tool_card_dialog_assets_and_api_are_wired(self):
        root = Path(__file__).resolve().parents[1]
        html = (root / "index.html").read_text(encoding="utf-8")
        script = (root / "panel_schedule_export_ui.js").read_text(encoding="utf-8")
        main = (root / "main.py").read_text(encoding="utf-8")

        self.assertIn('id="toolPanelScheduleExport"', html)
        self.assertIn('id="panelExportDlg"', html)
        self.assertIn('href="panel_schedule_export.css', html)
        self.assertIn('src="panel_schedule_export_ui.js', html)
        for element_id in re.findall(r'byId\("(\w+)"\)', script):
            self.assertIn(f'id="{element_id}"', html, element_id)
        self.assertIn("def export_panel_schedule_files(", main)
        self.assertIn("def save_panel_schedule_export(", main)
        self.assertIn("export_panel_schedule_files", script)
        self.assertIn("save_panel_schedule_export", script)

    def test_assets_are_packaged(self):
        root = Path(__file__).resolve().parents[1]
        for relative in ("ACIES Scheduler.spec", "build-config/ACIES Scheduler.spec", "build-config/build.ps1"):
            content = (root / relative).read_text(encoding="utf-8")
            self.assertIn("panel_schedule_export.css", content, relative)
            self.assertIn("panel_schedule_export_ui.js", content, relative)


class TestWorkbookRoundTrip(unittest.TestCase):
    def test_xlsx_export_writes_markdown_and_json(self):
        with tempfile.TemporaryDirectory() as folder:
            path = os.path.join(folder, "Panels.xlsx")
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "F"
            for address, value in panel_cells().items():
                sheet[address] = value
            hidden = book.create_sheet("Hidden")
            hidden.sheet_state = "hidden"
            hidden["A1"] = "nothing"
            book.save(path)

            out_dir = os.path.join(folder, "out")
            self.assertEqual(export.main([path, "-o", out_dir]), 0)
            written = sorted(os.listdir(out_dir))
            self.assertEqual(written, ["Panels panel schedules.json", "Panels panel schedules.md"])
            with open(os.path.join(out_dir, written[0]), encoding="utf-8") as handle:
                data = json.load(handle)
            self.assertEqual(data["panelCount"], 1)
            self.assertEqual(data["panels"][0]["circuits"][0]["description"], "RECEPTACLES")

    def test_missing_workbook_reports_error(self):
        self.assertEqual(export.main([os.path.join(tempfile.gettempdir(), "does-not-exist.xlsx")]), 1)


if __name__ == "__main__":
    unittest.main()
